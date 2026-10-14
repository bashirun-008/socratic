"""Knowledge base: turns slides, textbooks and lecture transcripts into source-tagged chunks
and retrieves the best ones with BM25 (pure Python, so no vector database to set up).

Every chunk keeps (source file, location) so answers can be cited, for example
"circuits_slides.pptx · slide 4" or "lecture.srt · 03:15".
"""
import io
import math
import re
from collections import Counter

STOP = set(
    "a an the and or of to in on for is are was were be been it this that these those with as by at "
    "from into than then so but if not no do does did can could should would will you your we they "
    "he she them his her its our their what which who whom how why when where about over under "
    "between".split()
)


def tokenize(text: str) -> list[str]:
    return [w for w in re.findall(r"[a-z0-9]+", text.lower()) if w not in STOP and len(w) > 1]


# ---------------- parsers: each returns [(location, text), ...] ----------------
def _pdf(data: bytes):
    from pypdf import PdfReader

    out = []
    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append((f"p. {i}", text))
    return out


def _pptx(data: bytes):
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text.strip():
            out.append((f"slide {i}", text))
    return out


_TS = re.compile(r"(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,]\d+\s*-->")


def _subs(text: str):
    """.srt / .vtt lecture transcripts -> ~60 second chunks tagged with a start time."""
    cues = []
    for block in re.split(r"\n\s*\n", text.replace("\r", "")):
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        for j, line in enumerate(lines):
            m = _TS.search(line)
            if m:
                h, mi, s = int(m.group(1) or 0), int(m.group(2)), int(m.group(3))
                body = " ".join(lines[j + 1:])
                body = re.sub(r"<[^>]+>", "", body)
                if body:
                    cues.append((h * 3600 + mi * 60 + s, body))
                break
    out, cur, start = [], [], 0
    for t, body in cues:
        if not cur:
            start = t
        cur.append(body)
        if t - start >= 60 or sum(len(c) for c in cur) > 700:
            out.append((_fmt(start), " ".join(cur)))
            cur = []
    if cur:
        out.append((_fmt(start), " ".join(cur)))
    return out


def _fmt(sec: int) -> str:
    h, m, s = sec // 3600, (sec % 3600) // 60, sec % 60
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _plain(text: str):
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    out, buf, n = [], "", 1
    for p in paras:
        if buf and len(buf) + len(p) > 800:
            out.append((f"section {n}", buf)); n += 1; buf = ""
        buf = (buf + "\n\n" + p).strip()
    if buf:
        out.append((f"section {n}", buf))
    return out


def _split(text: str, limit: int = 900):
    """Split long text on sentence boundaries into pieces of at most `limit` characters."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) <= limit:
        return [text]
    pieces, cur = [], ""
    for sent in re.split(r"(?<=[.!?])\s+|\n+", text):
        if cur and len(cur) + len(sent) + 1 > limit:
            pieces.append(cur.strip()); cur = ""
        cur += " " + sent
    if cur.strip():
        pieces.append(cur.strip())
    return pieces


class KnowledgeBase:
    def __init__(self):
        self.chunks: list[dict] = []
        self.files: dict[str, int] = {}
        self.by_id: dict[str, dict] = {}
        self._tok: list[Counter] = []
        self._df: Counter = Counter()
        self._avg = 1.0

    # ---- ingest ----
    def add_file(self, name: str, data: bytes) -> int:
        if name in self.files:
            return 0
        ext = name.lower().rsplit(".", 1)[-1]
        if ext == "pdf":
            segs = _pdf(data)
        elif ext == "pptx":
            segs = _pptx(data)
        elif ext in ("srt", "vtt"):
            segs = _subs(data.decode("utf-8", "ignore"))
        elif ext in ("txt", "md"):
            segs = _plain(data.decode("utf-8", "ignore"))
        else:
            raise ValueError(f"Unsupported file type: .{ext}")
        n = 0
        for loc, text in segs:
            for piece in _split(text):
                cid = f"S{len(self.chunks) + 1}"
                chunk = {"id": cid, "source": name, "loc": loc, "text": piece}
                self.chunks.append(chunk)
                self.by_id[cid] = chunk
                n += 1
        self.files[name] = n
        self._reindex()
        return n

    def _reindex(self):
        self._tok = [Counter(tokenize(c["text"])) for c in self.chunks]
        self._df = Counter()
        for t in self._tok:
            self._df.update(t.keys())
        self._avg = (sum(sum(t.values()) for t in self._tok) / max(len(self._tok), 1)) or 1.0

    # ---- retrieve ----
    def search(self, query: str, k: int = 3) -> list[dict]:
        q = set(tokenize(query))
        if not q or not self.chunks:
            return []
        N, k1, b = len(self.chunks), 1.5, 0.75
        scored = []
        for chunk, tf in zip(self.chunks, self._tok):
            dl, score = sum(tf.values()), 0.0
            for term in q:
                f = tf.get(term, 0)
                if not f:
                    continue
                idf = math.log(1 + (N - self._df[term] + 0.5) / (self._df[term] + 0.5))
                score += idf * f * (k1 + 1) / (f + k1 * (1 - b + b * dl / self._avg))
            if score > 0:
                scored.append((score, chunk))
        scored.sort(key=lambda x: -x[0])
        return [{**c, "score": round(s, 2)} for s, c in scored[:k]]

    def sample(self, k: int) -> list[dict]:
        if not self.chunks:
            return []
        step = max(len(self.chunks) // k, 1)
        return self.chunks[::step][:k]

    # ---- presentation ----
    def label(self, cid: str) -> str:
        if cid == "R1":
            return "Teacher notes"
        c = self.by_id.get(cid)
        return f"{c['source']} · {c['loc']}" if c else cid

    def format(self, hits: list[dict], limit: int = 700) -> str:
        return "\n".join(
            f"[{h['id']}] ({h['source']} · {h['loc']}) {h['text'][:limit]}" for h in hits
        )
